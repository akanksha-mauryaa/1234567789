"""
AEGIS MEDIA — Intelligent Media Classifier Lambda
====================================================
Triggered by S3 upload events.
Processes IMAGES (Rekognition) and DOCUMENTS (Textract / python-docx / direct read)
then runs deep NLP analysis (Comprehend) and saves results to DynamoDB.

Supported formats:
  Images : jpg, jpeg, png          → Rekognition labels + moderation
  PDF    : pdf                     → Textract text extraction → Comprehend NLP
  Word   : docx                    → python-docx text extraction → Comprehend NLP
  Text   : txt                     → Direct S3 read → Comprehend NLP

LAMBDA LAYER REQUIREMENT:
  For .docx support, add a Lambda Layer with `python-docx` installed.
  Build it like this:
    mkdir python && pip install python-docx -t python/
    zip -r python-docx-layer.zip python/
  Then upload as a Lambda Layer and attach to this function.
"""

import json
import urllib.parse
import boto3
import os
import io
import traceback
from datetime import datetime
from decimal import Decimal

s3 = boto3.client('s3')
rekognition = boto3.client('rekognition')
textract = boto3.client('textract')
comprehend = boto3.client('comprehend')
dynamodb = boto3.resource('dynamodb')

TABLE_NAME = os.environ.get('TABLE_NAME', 'media-processing-results')

# ── Safety keyword lists ──
UNSAFE_IMAGE_KEYWORDS = ['Gun', 'Firearm', 'Weapon', 'Violence', 'Blood', 'Nudity', 'Knife', 'Explosive']
UNSAFE_TEXT_KEYWORDS = ['bomb', 'kill', 'attack', 'terror', 'exploit', 'hack', 'malware', 'ransomware']
PII_ENTITY_TYPES = ['CREDIT_DEBIT_NUMBER', 'BANK_ACCOUNT_NUMBER', 'SSN', 'PIN', 'PASSWORD',
                    'CREDIT_DEBIT_CVV', 'CREDIT_DEBIT_EXPIRY', 'PHONE', 'EMAIL', 'ADDRESS',
                    'AWS_ACCESS_KEY', 'AWS_SECRET_KEY', 'IP_ADDRESS']

# ── Document Type Classification Patterns ──
# Each key is a document type, value is a list of indicator keywords/phrases.
# The classifier scores each type by counting how many indicators appear in the text.
DOCUMENT_TYPE_PATTERNS = {
    'Research Paper': [
        'abstract', 'introduction', 'methodology', 'literature review', 'conclusion',
        'references', 'et al', 'hypothesis', 'findings', 'doi', 'ieee', 'journal',
        'proceedings', 'citation', 'peer review', 'empirical', 'qualitative',
        'quantitative', 'research gap', 'fig.', 'table 1', 'experimental results',
        'related work', 'proposed method', 'state of the art', 'bibliography'
    ],
    'Resume / CV': [
        'experience', 'education', 'skills', 'objective', 'summary', 'certifications',
        'achievements', 'projects', 'responsibilities', 'proficient', 'bachelor',
        'master', 'university', 'gpa', 'linkedin', 'portfolio', 'references available',
        'work history', 'career objective', 'core competencies', 'professional summary'
    ],
    'Invoice': [
        'invoice', 'bill to', 'ship to', 'total', 'subtotal', 'tax', 'due date',
        'payment', 'quantity', 'unit price', 'amount due', 'balance due', 'purchase order',
        'invoice number', 'invoice date', 'remittance', 'net 30', 'gstin', 'gst'
    ],
    'Legal Document': [
        'whereas', 'hereby', 'herein', 'agreement', 'contract', 'clause', 'party',
        'jurisdiction', 'liability', 'indemnity', 'arbitration', 'governing law',
        'terms and conditions', 'confidentiality', 'non-disclosure', 'witness',
        'executed', 'binding', 'warrant', 'plaintiff', 'defendant', 'statute'
    ],
    'Business Report': [
        'executive summary', 'quarterly', 'annual report', 'revenue', 'profit',
        'growth', 'stakeholder', 'kpi', 'performance', 'forecast', 'budget',
        'roi', 'market analysis', 'swot', 'strategy', 'recommendations',
        'fiscal year', 'operating expenses', 'net income', 'dashboard'
    ],
    'Letter / Email': [
        'dear', 'sincerely', 'regards', 'to whom it may concern', 'yours faithfully',
        'attached', 'please find', 'looking forward', 'thank you for', 'best regards',
        'kind regards', 'yours truly', 'cc:', 'subject:', 're:'
    ],
    'Technical Documentation': [
        'api', 'endpoint', 'function', 'parameter', 'return value', 'example',
        'installation', 'configuration', 'usage', 'import', 'class', 'method',
        'deprecated', 'changelog', 'version', 'readme', 'sdk', 'documentation',
        'prerequisites', 'troubleshooting', 'syntax', 'command line'
    ],
    'Medical Report': [
        'patient', 'diagnosis', 'treatment', 'prescription', 'symptoms', 'clinical',
        'medical history', 'blood pressure', 'dosage', 'mg', 'prognosis',
        'laboratory', 'radiology', 'pathology', 'discharge summary', 'vital signs',
        'allergies', 'medication', 'follow-up', 'examination'
    ],
    'Educational Material': [
        'lesson', 'chapter', 'exercise', 'quiz', 'assignment', 'syllabus',
        'learning objectives', 'student', 'teacher', 'curriculum', 'exam',
        'grade', 'textbook', 'lecture', 'tutorial', 'semester', 'course',
        'homework', 'marks', 'question paper', 'answer key'
    ],
    'Presentation / Slides': [
        'slide', 'agenda', 'overview', 'key takeaways', 'next steps',
        'thank you', 'questions', 'outline', 'recap', 'demo',
        'presentation', 'deck', 'topics covered'
    ]
}


def lambda_handler(event, context):
    """Main entry point — triggered by S3 events."""
    table = dynamodb.Table(TABLE_NAME)

    for record in event.get('Records', []):
        bucket = record['s3']['bucket']['name']
        key = urllib.parse.unquote_plus(record['s3']['object']['key'], encoding='utf-8')

        file_name = key.split('/')[-1]
        if '_' in file_name:
            file_name = file_name.split('_', 1)[-1]

        file_ext = file_name.split('.')[-1].lower() if '.' in file_name else ''

        # --- Defaults ---
        file_type = 'unknown'
        rich_labels = []
        moderation_details = []
        is_safe = True
        status = 'skipped'
        doc_stats = {}  # Extra metadata for documents

        try:
            # ══════════════════════════════════════════
            # IMAGE PROCESSING (Rekognition)
            # ══════════════════════════════════════════
            if file_ext in ['jpg', 'jpeg', 'png']:
                file_type = 'image'

                # 1. Content Moderation
                mod_res = rekognition.detect_moderation_labels(
                    Image={'S3Object': {'Bucket': bucket, 'Name': key}},
                    MinConfidence=60
                )
                if mod_res.get('ModerationLabels'):
                    is_safe = False
                    moderation_details = [
                        {'name': m['Name'], 'confidence': Decimal(str(round(m['Confidence'], 1)))}
                        for m in mod_res['ModerationLabels']
                    ]

                # 2. Object / Scene Labels
                lab_res = rekognition.detect_labels(
                    Image={'S3Object': {'Bucket': bucket, 'Name': key}},
                    MaxLabels=15
                )
                for l in lab_res.get('Labels', []):
                    rich_labels.append({
                        'name': str(l['Name']),
                        'confidence': Decimal(str(round(l['Confidence'], 1)))
                    })

                # 3. Unsafe keyword check
                for l_obj in rich_labels:
                    if l_obj['name'] in UNSAFE_IMAGE_KEYWORDS:
                        is_safe = False
                        moderation_details.append({
                            'name': f"Security: {l_obj['name']}",
                            'confidence': l_obj['confidence']
                        })

                status = 'completed'

            # ══════════════════════════════════════════
            # DOCUMENT PROCESSING (Textract / python-docx / direct read)
            # ══════════════════════════════════════════
            elif file_ext in ['pdf', 'doc', 'docx', 'txt']:
                file_type = 'document'
                full_text = ""
                embedded_images = []

                # ── Step 1: Text Extraction (format-specific) ──
                if file_ext == 'pdf':
                    # PDF → Textract
                    text_res = textract.detect_document_text(
                        Document={'S3Object': {'Bucket': bucket, 'Name': key}}
                    )
                    for block in text_res.get('Blocks', []):
                        if block['BlockType'] == 'LINE':
                            full_text += block['Text'] + " "
                    print(f"[AEGIS] Textract extracted {len(full_text)} chars from PDF")

                elif file_ext == 'docx':
                    # DOCX → python-docx (requires Lambda Layer)
                    full_text = extract_docx_text(bucket, key)
                    print(f"[AEGIS] python-docx extracted {len(full_text)} chars from DOCX")

                    # ALSO extract embedded images from the docx file bytes directly!
                    try:
                        import zipfile
                        obj = s3.get_object(Bucket=bucket, Key=key)
                        file_bytes = obj['Body'].read()
                        with zipfile.ZipFile(io.BytesIO(file_bytes)) as z:
                            image_files = [f for f in z.namelist() if f.startswith('word/media/')]
                            for img_name in image_files:
                                img_data = z.read(img_name)
                                try:
                                    rek_res = rekognition.detect_moderation_labels(
                                        Image={'Bytes': img_data},
                                        MinConfidence=60
                                    )
                                    mod_labels = rek_res.get('ModerationLabels', [])
                                    is_img_safe = len(mod_labels) == 0
                                    
                                    img_alerts = [
                                        {'name': str(m['Name']), 'confidence': Decimal(str(round(m['Confidence'], 1)))}
                                        for m in mod_labels
                                    ]
                                    
                                    embedded_images.append({
                                        'name': str(img_name.split('/')[-1]),
                                        'is_safe': is_img_safe,
                                        'alerts': img_alerts
                                    })
                                    
                                    if not is_img_safe:
                                        is_safe = False
                                        # Add embedded image alert to moderation details
                                        for alert in img_alerts:
                                            moderation_details.append({
                                                'name': f"Unsafe Embedded Image ({img_name.split('/')[-1]}): {alert['name']}",
                                                'confidence': alert['confidence']
                                            })
                                except Exception as img_err:
                                    print(f"[AEGIS] Failed to scan embedded image {img_name}: {img_err}")
                    except Exception as zip_err:
                        print(f"[AEGIS] Failed to parse zip/docx embedded images: {zip_err}")

                elif file_ext == 'txt':
                    # TXT → Direct S3 read
                    obj = s3.get_object(Bucket=bucket, Key=key)
                    full_text = obj['Body'].read().decode('utf-8', errors='replace')
                    print(f"[AEGIS] Direct read: {len(full_text)} chars from TXT")

                elif file_ext == 'doc':
                    # .doc (legacy binary) — not easily parseable in Lambda
                    # Mark as received; we can't reliably extract text from .doc
                    rich_labels.append({
                        'name': 'DOC File Received (Legacy Format)',
                        'confidence': Decimal('95.0')
                    })
                    print(f"[AEGIS] .doc is a legacy binary format — skipping text extraction")

                # ── Step 2: Document Type Classification ──
                if full_text.strip():
                    doc_type, doc_type_conf = classify_document_type(full_text)
                    # Insert as the FIRST label so it's most prominent
                    rich_labels.insert(0, {
                        'name': f'DocType: {doc_type}',
                        'confidence': Decimal(str(doc_type_conf))
                    })
                    print(f"[AEGIS] Document classified as: {doc_type} ({doc_type_conf}% confidence)")

                # ── Step 3: Deep NLP Analysis ──
                if full_text.strip():
                    nlp_results = run_deep_nlp(full_text)
                    rich_labels.extend(nlp_results['labels'])
                    moderation_details.extend(nlp_results['moderation'])
                    doc_stats = nlp_results['stats']
                    # Store the document type in stats too
                    if full_text.strip():
                        doc_stats['document_type'] = doc_type
                        doc_stats['document_type_confidence'] = Decimal(str(doc_type_conf))
                    if nlp_results['has_pii']:
                        is_safe = False

                # ── Fallback label ──
                if not rich_labels:
                    rich_labels = [{'name': 'Document Analyzed', 'confidence': Decimal('100.0')}]

                status = 'completed'

        except Exception as e:
            print(f"[AEGIS] Processing error for {file_name}: {str(e)}")
            print(traceback.format_exc())
            rich_labels = [{'name': f'Processing Error: {type(e).__name__}', 'confidence': Decimal('0.0')}]
            status = 'error'

        # --- Save to DynamoDB ---
        timestamp = datetime.utcnow().strftime('%Y-%m-%d %H:%M')

        item = {
            'file_id': key,
            'file_name': file_name,
            'file_type': file_type,
            'labels': rich_labels,
            'moderation_details': moderation_details,
            'label_count': len(rich_labels),
            'is_safe': is_safe,
            'status': status,
            'timestamp': timestamp
        }

        # Add document stats if available
        if doc_stats:
            item['doc_stats'] = doc_stats

        # Save embedded images to DynamoDB if any were detected in DOCX
        try:
            if 'embedded_images' in locals() and embedded_images:
                item['embedded_images'] = embedded_images
        except Exception:
            pass

        print(f"[AEGIS] Saving item with {len(rich_labels)} labels, safe={is_safe}, status={status}")
        table.put_item(Item=item)

    return {'statusCode': 200, 'body': 'Success'}


# ══════════════════════════════════════════════════════════
# HELPER: Extract text from .docx using python-docx
# ══════════════════════════════════════════════════════════
def extract_docx_text(bucket, key):
    """Download .docx from S3 and extract all paragraph text using python-docx."""
    try:
        from docx import Document as DocxDocument
    except ImportError:
        print("[AEGIS] python-docx not available — add it as a Lambda Layer!")
        print("[AEGIS] Falling back to S3 raw read...")
        # Fallback: try to read raw bytes and extract printable strings
        obj = s3.get_object(Bucket=bucket, Key=key)
        raw = obj['Body'].read()
        # Extract printable ASCII sequences (rough fallback)
        import re
        text_chunks = re.findall(rb'[\x20-\x7E]{10,}', raw)
        return ' '.join(chunk.decode('ascii', errors='ignore') for chunk in text_chunks[:100])

    # Download file into memory
    obj = s3.get_object(Bucket=bucket, Key=key)
    file_bytes = obj['Body'].read()

    # Parse with python-docx
    doc = DocxDocument(io.BytesIO(file_bytes))

    paragraphs = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            paragraphs.append(text)

    # Also extract text from tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    paragraphs.append(text)

    return ' '.join(paragraphs)


# ══════════════════════════════════════════════════════════
# HELPER: Deep NLP analysis using AWS Comprehend
# ══════════════════════════════════════════════════════════
def run_deep_nlp(full_text):
    """
    Run comprehensive NLP analysis on extracted text.
    Returns dict with: labels, moderation, stats, has_pii
    """
    labels = []
    moderation = []
    has_pii = False
    text_snippet = full_text[:5000]  # Comprehend limit per call

    # ── 1. Language Detection ──
    try:
        lang_res = comprehend.detect_dominant_language(Text=text_snippet[:1000])
        top_lang = lang_res.get('Languages', [{}])[0]
        lang_code = top_lang.get('LanguageCode', 'unknown')
        lang_score = top_lang.get('Score', 0)
        labels.append({
            'name': f'Language: {lang_code.upper()}',
            'confidence': Decimal(str(round(lang_score * 100, 1)))
        })
    except Exception as e:
        print(f"[AEGIS] Language detection failed: {e}")

    # ── 2. Sentiment Analysis ──
    try:
        sent_res = comprehend.detect_sentiment(Text=text_snippet, LanguageCode='en')
        sentiment = sent_res.get('Sentiment', 'UNKNOWN')
        scores = sent_res.get('SentimentScore', {})
        # Add the dominant sentiment
        sentiment_conf = scores.get(sentiment.capitalize(), scores.get('Neutral', 0))
        labels.append({
            'name': f'Sentiment: {sentiment.capitalize()}',
            'confidence': Decimal(str(round(sentiment_conf * 100, 1)))
        })
    except Exception as e:
        print(f"[AEGIS] Sentiment analysis failed: {e}")

    # ── 3. Key Phrases ──
    try:
        kp_res = comprehend.detect_key_phrases(Text=text_snippet, LanguageCode='en')
        key_phrases = kp_res.get('KeyPhrases', [])
        # Take top 8 by score
        key_phrases.sort(key=lambda x: x.get('Score', 0), reverse=True)
        for kp in key_phrases[:8]:
            phrase_text = kp['Text'][:50]  # Truncate long phrases
            labels.append({
                'name': f'Topic: {phrase_text}',
                'confidence': Decimal(str(round(kp['Score'] * 100, 1)))
            })
    except Exception as e:
        print(f"[AEGIS] Key phrases failed: {e}")

    # ── 4. Entity Detection ──
    try:
        ent_res = comprehend.detect_entities(Text=text_snippet, LanguageCode='en')
        entities = ent_res.get('Entities', [])
        # Deduplicate by (type, text) and take top entities
        seen_entities = set()
        for ent in sorted(entities, key=lambda x: x['Score'], reverse=True):
            if ent['Score'] > 0.8:
                ent_key = (ent['Type'], ent['Text'][:30])
                if ent_key not in seen_entities and len(seen_entities) < 10:
                    seen_entities.add(ent_key)
                    labels.append({
                        'name': f"{ent['Type']}: {ent['Text'][:40]}",
                        'confidence': Decimal(str(round(ent['Score'] * 100, 1)))
                    })
    except Exception as e:
        print(f"[AEGIS] Entity detection failed: {e}")

    # ── 5. PII Detection ──
    try:
        pii_res = comprehend.detect_pii_entities(Text=text_snippet, LanguageCode='en')
        pii_entities = pii_res.get('Entities', [])
        pii_found = []
        for pii in pii_entities:
            if pii['Score'] > 0.7 and pii['Type'] in PII_ENTITY_TYPES:
                pii_found.append(pii)
                moderation.append({
                    'name': f"PII: {pii['Type'].replace('_', ' ').title()}",
                    'confidence': Decimal(str(round(pii['Score'] * 100, 1)))
                })
        if pii_found:
            has_pii = True
            labels.append({
                'name': f'⚠ PII Detected ({len(pii_found)} items)',
                'confidence': Decimal('100.0')
            })
    except Exception as e:
        print(f"[AEGIS] PII detection failed: {e}")

    # ── 6. Unsafe Text Keyword Scan ──
    text_lower = full_text.lower()
    for keyword in UNSAFE_TEXT_KEYWORDS:
        if keyword in text_lower:
            moderation.append({
                'name': f'Flagged Keyword: {keyword}',
                'confidence': Decimal('85.0')
            })

    # ── 7. Document Statistics ──
    words = full_text.split()
    word_count = len(words)
    char_count = len(full_text)
    sentence_count = full_text.count('.') + full_text.count('!') + full_text.count('?')

    stats = {
        'word_count': word_count,
        'char_count': char_count,
        'sentence_count': max(sentence_count, 1),
        'avg_word_length': Decimal(str(round(char_count / max(word_count, 1), 1)))
    }

    labels.append({
        'name': f'Words: {word_count:,}',
        'confidence': Decimal('100.0')
    })

    return {
        'labels': labels,
        'moderation': moderation,
        'stats': stats,
        'has_pii': has_pii
    }


# ══════════════════════════════════════════════════════════
# HELPER: Classify document type from text content
# ══════════════════════════════════════════════════════════
def classify_document_type(text):
    """
    Classify a document into a category (Research Paper, Resume, Invoice, etc.)
    by scoring keyword pattern matches against the extracted text.
    
    Returns: (document_type: str, confidence: float)
    """
    text_lower = text.lower()
    scores = {}

    for doc_type, keywords in DOCUMENT_TYPE_PATTERNS.items():
        matched = 0
        for kw in keywords:
            if kw.lower() in text_lower:
                matched += 1
        if matched > 0:
            # Confidence = percentage of keywords matched, scaled to 60-99 range
            raw_score = matched / len(keywords)
            confidence = round(60 + (raw_score * 39), 1)  # Maps 0→60%, 1→99%
            scores[doc_type] = confidence

    if not scores:
        return ('General Document', 70.0)

    # Sort by confidence descending
    ranked = sorted(scores.items(), key=lambda x: x[1], reverse=True)
    best_type, best_conf = ranked[0]

    # If top two are very close (within 5%), mark it as a hybrid
    if len(ranked) >= 2 and (ranked[0][1] - ranked[1][1]) < 5:
        second_type = ranked[1][0]
        best_type = f"{best_type} / {second_type}"
        # Average the two confidences
        best_conf = round((ranked[0][1] + ranked[1][1]) / 2, 1)

    print(f"[AEGIS] DocType scores: {dict(ranked[:5])}")
    return (best_type, best_conf)
